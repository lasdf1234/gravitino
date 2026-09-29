#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate Chinese TMS design docx from markdown source."""

import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


NUM_LIST_RE = re.compile(r"^(\d+)\.\s+(.*)$")
BULLET_RE = re.compile(r"^-\s+(.*)$")
CHECKBOX_RE = re.compile(r"^- \[( |x)\]\s+(.*)$")


def set_doc_font(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "PingFang SC"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")


def add_rich_text(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def add_body_paragraph(doc: Document, text: str, indent: bool = False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
    add_rich_text(p, text)
    return p


def append_body_text(paragraph, text: str) -> None:
    paragraph.add_run(" ")
    add_rich_text(paragraph, text.strip())


def is_continuation_line(raw_line: str) -> bool:
    if not raw_line.startswith((" ", "\t")):
        return False
    stripped = raw_line.strip()
    if not stripped:
        return False
    if stripped.startswith(("#", "|", "```", "- ", "- [", ">")):
        return False
    if NUM_LIST_RE.match(stripped):
        return False
    return True


def add_code_block(doc: Document, text: str) -> None:
    for line in text.splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        run = p.add_run(line)
        run.font.name = "Menlo"
        run.font.size = Pt(9)
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Menlo")


def parse_table_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def is_table_sep(line: str) -> bool:
    return bool(re.match(r"^\|[\s\-:|]+\|\s*$", line.strip()))


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            cell_text = row[j] if j < len(row) else ""
            table.rows[i].cells[j].text = cell_text


def md_to_docx(md_path: Path, out_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    set_doc_font(doc)

    i = 0
    in_code = False
    code_buf: list[str] = []
    table_buf: list[list[str]] = []
    last_paragraph = None
    last_is_list_item = False

    def flush_list_context() -> None:
        nonlocal last_paragraph, last_is_list_item
        last_paragraph = None
        last_is_list_item = False

    while i < len(lines):
        raw_line = lines[i]
        line = raw_line.rstrip()

        if line.strip().startswith("<!--"):
            flush_list_context()
            while i < len(lines) and "-->" not in lines[i]:
                i += 1
            i += 1
            continue

        if line.strip().startswith("```"):
            flush_list_context()
            if in_code:
                add_code_block(doc, "\n".join(code_buf))
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if line.strip().startswith("|") and "|" in line.strip()[1:]:
            flush_list_context()
            if not is_table_sep(line):
                table_buf.append(parse_table_row(line))
            i += 1
            if i >= len(lines) or not lines[i].strip().startswith("|"):
                add_table(doc, table_buf)
                table_buf = []
            continue

        if table_buf:
            flush_list_context()
            add_table(doc, table_buf)
            table_buf = []

        stripped = line.strip()
        if not stripped:
            flush_list_context()
            i += 1
            continue

        if stripped == "---":
            flush_list_context()
            doc.add_paragraph()
            i += 1
            continue

        if is_continuation_line(raw_line) and last_is_list_item and last_paragraph is not None:
            append_body_text(last_paragraph, stripped)
            i += 1
            continue

        if stripped.startswith("#"):
            flush_list_context()
            level = len(stripped) - len(stripped.lstrip("#"))
            title = stripped[level:].strip()
            doc.add_heading(title, level=min(level, 4))
            i += 1
            continue

        checkbox_match = CHECKBOX_RE.match(stripped)
        if checkbox_match:
            mark = "☑" if checkbox_match.group(1) == "x" else "☐"
            text = f"{mark} {checkbox_match.group(2)}"
            last_paragraph = add_body_paragraph(doc, text, indent=True)
            last_is_list_item = True
            i += 1
            continue

        bullet_match = BULLET_RE.match(stripped)
        if bullet_match:
            text = f"• {bullet_match.group(1)}"
            last_paragraph = add_body_paragraph(doc, text, indent=True)
            last_is_list_item = True
            i += 1
            continue

        num_match = NUM_LIST_RE.match(stripped)
        if num_match:
            number, body = num_match.groups()
            text = f"{number}. {body}"
            last_paragraph = add_body_paragraph(doc, text, indent=True)
            last_is_list_item = True
            i += 1
            continue

        flush_list_context()
        if stripped.startswith(">"):
            last_paragraph = add_body_paragraph(doc, stripped[1:].strip())
        else:
            last_paragraph = add_body_paragraph(doc, stripped)
        last_is_list_item = False
        i += 1

    if table_buf:
        add_table(doc, table_buf)
    if in_code and code_buf:
        add_code_block(doc, "\n".join(code_buf))

    doc.add_paragraph()
    p = doc.add_paragraph("— 文档结束 —")
    p.runs[0].italic = True

    doc.save(out_path)
    print(f"Wrote {out_path}")


if __name__ == "__main__":
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).parent / "tms-design-zh.md"
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else Path.home() / "Desktop" / "TMS设计文档-中文翻译.docx"
    md_to_docx(src, dst)
